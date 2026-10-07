"""Build the manuscript DOCX from its reviewable Markdown sources, offline.

Tables follow academic conventions: a title paragraph above, then a table with a top rule, a rule under the header
and a bottom rule, no vertical lines and no shading. Figures are followed by their caption paragraph.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "medrxiv-manuscript.docx"
TEXT_WIDTH = 6.9
# Relative column widths, keyed by the table label in its title paragraph.
WIDTHS = {
    "Table 1": [0.6, 1.95, 1.35, 1.35, 0.5, 0.75, 0.6],
    "Table 2": [1.7, 1.45, 0.85, 1.45, 0.85],
    "Table 3": [0.6, 1.8, 1.2, 1.25, 0.62, 0.52, 0.6, 0.5],
    "Table 4": [2.0, 0.8, 0.8, 0.85, 0.85, 0.8, 0.8],
    "Supplementary Table 1": [1.45, 1.1, 1.35, 1.35, 1.25, 0.95],
    "Supplementary Table 2": [0.8, 2.6, 1.4, 1.2],
    "Supplementary Table 3": [0.7, 1.3, 1.0, 1.3, 1.3, 1.3],
    "Supplementary Table 4": [0.7, 3.6, 1.6],
    "Supplementary Table 5": [0.6, 1.6, 0.6, 0.7, 0.75, 0.85, 0.7, 0.6],
    "Supplementary Table 6": [0.6, 1.2, 1.0, 0.7, 0.8, 0.7, 1.5],
    "Supplementary Table 7": [2.0, 0.8, 0.8, 0.85, 0.85, 0.8, 0.8],
    "Supplementary Table 8": [1.6, 0.8, 0.8, 0.85, 0.85, 0.8, 0.8],
    "Supplementary Table 9": [1.3, 1.35, 1.1, 1.5, 2.2],
    "Supplementary Table 10": [1.6, 0.7, 0.7, 0.8, 0.9, 0.8, 0.7],
}
LABEL = re.compile(r"^\*\*((?:Supplementary )?Table \d+)\.")


def inline(paragraph, text: str, size: float | None = None) -> None:
    # Preserve citation numbers and URLs as readable text.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for piece in re.split(r"(\*\*.*?\*\*|`.*?`)", text):
        bold = piece.startswith("**") and piece.endswith("**")
        code = piece.startswith("`") and piece.endswith("`")
        run = paragraph.add_run(piece[2:-2] if bold else piece[1:-1] if code else piece)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Batang")
        if size:
            run.font.size = Pt(size)


def border(element, tag: str, sides: dict[str, int]) -> None:
    """Attach single-line borders of the given eighth-point sizes; size 0 removes the side."""
    box = OxmlElement(tag)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if side not in sides:
            continue
        e = OxmlElement(f"w:{side}")
        size = sides[side]
        e.set(qn("w:val"), "single" if size else "nil")
        if size:
            e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), "000000"); e.set(qn("w:space"), "0")
        box.append(e)
    element.append(box)


def table(doc, lines: list[str], label: str | None) -> None:
    raw = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    spec = next(r for r in raw if all(re.fullmatch(r":?-+:?", c) for c in r))
    right = [c.endswith(":") for c in spec]
    rows = [r for r in raw if r is not spec]
    n = len(rows[0])
    widths = WIDTHS.get(label, [TEXT_WIDTH / n] * n)
    if len(widths) != n:
        raise ValueError(f"Width rule for {label} has {len(widths)} columns, table has {n}")
    widths = [w * TEXT_WIDTH / sum(widths) for w in widths]
    tab = doc.add_table(rows=0, cols=n)
    tab.autofit = False
    border(tab._tbl.tblPr, "w:tblBorders", {"top": 8, "bottom": 8, "left": 0, "right": 0, "insideH": 0, "insideV": 0})
    for column, width in zip(tab.columns, widths):
        column.width = Inches(width)
    size = 9 if n >= 6 else 9.5
    for i, values in enumerate(rows):
        if len(values) != n:
            raise ValueError(f"Malformed table row {values}")
        row = tab.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if i == 0:
            repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); props.append(repeat)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM if i == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, amount in (("top", 30), ("left", 70), ("bottom", 30), ("right", 70)):
                e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(amount)); e.set(qn("w:type"), "dxa"); margins.append(e)
            tcpr.append(margins)
            if i == 0:
                border(tcpr, "w:tcBorders", {"bottom": 6})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = i < len(rows) - 1 and len(rows) <= 14
            if right[j] and j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            mark = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); mark.append(sz)
            p._p.get_or_add_pPr().append(mark)  # Empty cells otherwise take the 11 pt body size and enlarge the row.
            inline(p, value, size)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def next_content(lines: list[str], i: int) -> str:
    j = i + 1
    while j < len(lines) and not lines[j].strip():
        j += 1
    return lines[j].strip() if j < len(lines) else ""


def add_markdown(doc, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    i, label = 0, None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            table(doc, block, label)
            continue
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = True
                r = p.add_run(lines[i])
                r.font.name = "Consolas"
                r.font.size = Pt(9.5 if re.search(r"[가-힣]", lines[i]) else 9)
                r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Batang")
                i += 1
            i += 1
            continue
        if line.startswith("#"):
            marks, title = line.split(" ", 1)
            if len(marks) == 1 and path.name == "manuscript.md":
                doc.add_paragraph(title, "Title")
            else:
                p = doc.add_paragraph(title, "Heading 1" if len(marks) <= 2 else "Heading 2")
                if len(marks) == 1:
                    p.paragraph_format.page_break_before = True
        elif line.startswith("!["):
            match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
            assert match
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            shape = p.add_run().add_picture(str((path.parent / match.group(2)).resolve()), width=Inches(TEXT_WIDTH))
            shape._inline.docPr.set("descr", match.group(1))
        else:
            p = doc.add_paragraph()
            found = LABEL.match(line)
            if found:
                label = found.group(1)
            caption = line.startswith(("**Table", "**Supplementary Table", "**Figure"))
            if caption:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(6)
                if line.startswith(("**Table", "**Supplementary Table")):
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.keep_with_next = True
                    # Long tables that fit one page start on a new page, so grouped labels are never split from their rows.
                    j = i + 1
                    while j < len(lines) and not lines[j].strip().startswith("|"):
                        j += 1
                    rows = 0
                    while j < len(lines) and lines[j].strip().startswith("|"):
                        rows += 1; j += 1
                    if 16 <= rows <= 40:
                        p.paragraph_format.page_break_before = True
                else:
                    p.paragraph_format.space_after = Pt(10)
            if next_content(lines, i).startswith("|"):
                p.paragraph_format.keep_with_next = True
            if re.match(r"^(\d+|S\d+)\. ", line):
                p.paragraph_format.space_after = Pt(6)
            if line.startswith("- "):
                p.style = "List Bullet"
                line = line[2:]
            inline(p, line, 10 if caption else None)
        i += 1


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.footer_distance = Inches(0.35)
    for name, size in [("Normal", 11), ("Title", 18), ("Heading 1", 14), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Batang")
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.1
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
        rpr = style._element.get_or_add_rPr()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rpr.rFonts.attrib.pop(qn("w:" + attr), None)
        for tag in ("w:spacing", "w:kern"):
            for el in rpr.findall(qn(tag)):
                rpr.remove(el)
    title = doc.styles["Title"]
    for el in title._element.get_or_add_pPr().findall(qn("w:pBdr")):
        el.getparent().remove(el)
    title.font.bold = True
    title.font.size = Pt(16)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.15
    doc.styles["Normal"].paragraph_format.widow_control = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    doc.core_properties.title = "Referring uncertain answers by decision-model confidence: Jev, OpenAI Decisions and Clef on Korean and US medical licensing examination questions"
    doc.core_properties.author = ""
    doc.core_properties.subject = "Exploratory in silico study for medRxiv"
    sources = [HERE / "manuscript.md", HERE / "supplementary.md"]
    for path in sources:
        add_markdown(doc, path)
    doc.save(OUT)
    print(json.dumps({"output": str(OUT), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "images": len(doc.inline_shapes),
                      "source_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}}, indent=2))


if __name__ == "__main__":
    main()
